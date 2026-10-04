"""End-to-end test of Step 5 on a small synthetic artifact folder shaped like the Step 1–4 outputs."""
import json
import re
import tempfile
import unittest
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import step5_report  # adds the repo root (eval.py) to sys.path
import eval as ev
from eval import save_predictions

N_VAL, N_TEST = 90, 99


def _probs(y, rng, strength):
    logits = rng.normal(size=(len(y), 9)); logits[np.arange(len(y)), y] += strength
    e = np.exp(logits - logits.max(1, keepdims=True)); return e / e.sum(1, keepdims=True), logits


def _run(root, exp_id, seed, tag, val_f1, rng, y_val, strength, names_val, label=None):
    d = root / "runs" / exp_id / f"seed{seed}"; d.mkdir(parents=True)
    hist = pd.DataFrame({"epoch": [1, 2, 3], "train_loss": [1.2, .6, .4], "lr": [1e-3, 5e-4, 1e-5], "train_seconds": [10, 10, 10],
                         "val_loss": [.9, .5, .55], "val_macro_f1": [.6, val_f1, val_f1 - .01], "val_top1": [.7, .9, .89]})
    hist.to_csv(d / "history.csv", index=False)
    probs, logits = _probs(y_val, rng, strength)
    save_predictions(root / "predictions" / f"{exp_id}_seed{seed}_val.csv", names_val, y_val, probs)
    np.save(d / "val_logits.npy", logits)
    f1 = ev.compute_metrics(y_val, probs.argmax(1), probs)["macro_f1"]
    curve = root / "curves" / f"{exp_id}_{label or tag.split('.')[0]}.png"
    summary = {"exp_id": exp_id, "seed": seed, "val_macro_f1": f1, "val_top1": .9, "pretrained_tag": tag,
               "backbone": tag.split(".")[0], "curve": str(curve), "history": str(d / "history.csv"), "best_epoch": 2,
               "val_f1_per_class": [f1] * 9, "mean_train_seconds": 10.0, "params_m": 25.0, "gmac": 4.1}
    (d / "summary.json").write_text(json.dumps(summary)); (d / "config.json").write_text(json.dumps(
        {"backbone": tag, "epochs": 3, "batch_size": 64, "grad_accum_steps": 1, "init": "finetune", "aug": "basic", "loss": "ce"}))
    (d / "split_report.json").write_text(json.dumps({"n": {"train": 10501, "val": 3501, "test": 3507}}))
    return summary


def make_artifacts(root):
    rng = np.random.default_rng(0); labels = root / "labels"; labels.mkdir(parents=True); (root / "predictions").mkdir()
    y_val = np.arange(N_VAL) % 9; y_test = np.arange(N_TEST) % 9
    names_val = [f"v{i}.jpg" for i in range(N_VAL)]; names_test = [f"t{i}.jpg" for i in range(N_TEST)]
    pd.DataFrame({"Filename": names_val, "Label": y_val}).to_csv(labels / "val_subset0.csv", index=False)
    pd.DataFrame({"Filename": names_test, "Label": y_test}).to_csv(labels / "test_subset0.csv", index=False)
    pd.DataFrame({"Filename": names_test, "Label": y_test, "Species": [ev.CLASS_NAMES[i] for i in y_test]}).to_csv(labels / "labels.csv", index=False)
    tags = ["resnet50.a1_in1k", "convnext_tiny.fb_in1k", "deit_small_patch16_224.fb_in1k", "efficientnet_b0.ra_in1k", "mobilenetv3_large_100.ra_in1k"]
    bb = []
    for i, tag in enumerate(tags, 1):
        s = _run(root, f"B0{i}", 0, tag, .9, rng, y_val, 2 + i * .3, names_val)
        bb.append({"exp_id": f"B0{i}", "backbone": tag.split(".")[0], "pretrained_tag": tag, "params_m": 5.0 * i, "gmac": .4 * i,
                   "img_size": 224, "epochs": 3, "seed": 0, "physical_batch": 64, "effective_batch": 64, "val_macro_f1": s["val_macro_f1"],
                   "val_top1": .9, "best_epoch": 2, "train_seconds_epoch": 10.0 * i, "latency_p50_ms": 3.0 * i, "latency_p95_ms": 3.5 * i,
                   "latency_dtype": "amp", "gpu": "Tesla T4", "curve": s["curve"], "note": "x", "overfit_flag": False})
    pd.DataFrame(bb).to_csv(root / "backbones.csv", index=False)
    best = pd.DataFrame(bb).sort_values("val_macro_f1").iloc[-1]
    (root / "backbone_selection.json").write_text(json.dumps({"best_quality": best.exp_id, "best_quality_backbone": best.pretrained_tag,
                                                               "balanced": "B04", "balanced_backbone": tags[3]}))
    tr = []
    for seed in range(3):
        s = _run(root, "T00", seed, best.pretrained_tag, .9, rng, y_val, 3, names_val, f"baseline_seed{seed}")
        tr.append(("T00", "Baseline", "T00", seed, s))
    for k, (axis, change) in enumerate([("A", "frozen backbone"), ("B", "ColorJitter"), ("C", "label smoothing 0.1"), ("Combined", "T02 + T03")], 1):
        s = _run(root, f"T{k:02d}", 0, best.pretrained_tag, .9, rng, y_val, 2.5 + k * .3, names_val, change.replace(" ", "_"))
        tr.append((f"T{k:02d}", axis, change, 0, s))
    std = float(np.std([x[4]["val_macro_f1"] for x in tr[:3]], ddof=1)); base0 = tr[0][4]["val_macro_f1"]
    pd.DataFrame([{"exp_id": e, "backbone": best.pretrained_tag, "axis": a, "change_vs_T00": c, "seed": sd, "val_macro_f1": s["val_macro_f1"],
                   "val_top1": .9, "delta_vs_T00_seed0": None if e == "T00" else s["val_macro_f1"] - base0, "baseline_macro_f1_std": std,
                   "verdict_vs_noise": "baseline" if e == "T00" else "phân biệt được", "chinee_apple_f1": .8, "snake_weed_f1": .8,
                   "best_epoch": 2, "train_seconds_epoch": 10.0, "curve": s["curve"]} for e, a, c, sd, s in tr]).to_csv(root / "training.csv", index=False)
    (root / "training_selection.json").write_text(json.dumps({"backbone": best.pretrained_tag, "exp_id": "T03", "overrides": {"loss": "ls"},
                                                              "val_macro_f1": tr[5][4]["val_macro_f1"], "combo_components": ["T02", "T03"],
                                                              "combo_interaction": "cộng dồn một phần"}))
    t03 = ev.read_pred(str(root / "predictions" / "T03_seed0_val.csv"))
    inf = []; lat = []
    for code, method, strength in (("I00", "1-view FP32", None), ("I01", "TTA flip", 4.0), ("I07", "Temperature scaling", None)):
        probs = t03.probs if strength is None else _probs(y_val, rng, strength)[0]
        save_predictions(root / "predictions" / f"{code}_seed0_val.csv", names_val, y_val, probs)
        m = ev.compute_metrics(y_val, probs.argmax(1), probs)
        inf.append({"exp_id": code, "method": method, "checkpoint": "T03", "K": 1, "val_macro_f1": m["macro_f1"], "val_top1": m["top1"],
                    "ece": m["ece"], "latency_p50_ms": 4.0, "latency_p95_ms": 5.0, "latency_p99_ms": 6.0, "dtype": "fp32", "note": "", "relative_cost": 1.0})
        for batch in (1, 32):
            lat.append({"batch": batch, "p50_ms": 4.0 * batch ** .5, "p95_ms": 5.0, "p99_ms": 6.0, "mean_ms": 4.1, "images_per_s": 250.0,
                        "dtype": "fp32", "gpu": "Tesla T4", "img_size": 224, "warmup": 10, "iterations": 50, "torch": "2.x",
                        "includes_preprocessing": False, "exp_id": code, "method": method, "fused_bn": False, "measurement_scope": "gpu"})
    pd.DataFrame(inf).to_csv(root / "inference.csv", index=False); pd.DataFrame(lat).to_csv(root / "latency.csv", index=False)
    (root / "inference_selection.json").write_text(json.dumps({"realtime_method": "I07", "offline_method": "I01"}))
    (root / "final_lock.json").write_text(json.dumps({"backbone": best.pretrained_tag, "training_source": "T03",
                                                      "training_overrides": {"loss": "ls"}, "seeds": [0, 1, 2], "epochs": 3, "batch_size": 64}))
    for seed in range(3):
        _run(root, "F01", seed, best.pretrained_tag, .92, rng, y_val, 3.5, names_val, f"seed{seed}")
        for tag, strength in (("F01", 3.6), ("F01uncal", 3.6), ("T00", 2.8)):
            save_predictions(root / "predictions" / f"{tag}_seed{seed}_test.csv", names_test, y_test, _probs(y_test, rng, strength)[0])
    save_predictions(root / "predictions" / "F01uncal_seed0_val.csv", names_val, y_val, _probs(y_val, rng, 3)[0])
    common = ["--test-csv", str(labels / "test_subset0.csv"), "--labels", str(labels / "labels.csv")]
    with redirect_stdout(StringIO()):
        for tag in ("F01", "T00"):
            assert ev.main(["score", "--pred", str(root / "predictions" / f"{tag}_seed*_test.csv"), *common, "--tag", tag, "--out", str(root / "eval_out")]) == 0
        assert ev.main(["grade", "--final", str(root / "predictions" / "F01_seed*_test.csv"), "--baseline", str(root / "predictions" / "T00_seed*_test.csv"),
                        "--uncal", str(root / "predictions" / "F01uncal_seed*_test.csv"), "--latency-p95-ms", "5", *common, "--out", str(root / "eval_out")]) == 0
    return labels


class TestStep5(unittest.TestCase):
    def test_build_deliverables_end_to_end(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "artifacts"; sub = Path(tmp) / "submission"; labels = make_artifacts(root)
            with redirect_stdout(StringIO()):
                out = step5_report.build_deliverables(root, sub, labels)
            wb = load_workbook(out["results_xlsx"])
            for sheet in ("Backbones", "Training", "Inference", "Final", "PerClass", "Latency", "Summary"):
                self.assertIn(sheet, wb.sheetnames)
            checks = out["checks"].set_index("check")
            self.assertTrue(checks.loc["Mọi lần huấn luyện B/T/F có ảnh curves/", "ok"])
            self.assertTrue(checks.loc["Final F01 khớp eval_out/F01_per_seed.csv (eval.py score)", "ok"])
            self.assertTrue(checks.loc["Final T00 mean khớp eval_out/T00_summary.json", "ok"])
            self.assertTrue(checks.loc["Inference I00 = Training T03 seed 0", "ok"])
            # 5 B + 3 T00 + 4 T + 3 F01 = 15 curves
            self.assertEqual(len(list((sub / "curves").glob("*.png"))), 15)
            self.assertTrue((sub / "curves" / "B01_resnet50.png").is_file())
            self.assertEqual(len(list((sub / "predictions").glob("F01_seed*_test.csv"))), 3)
            report = (sub / "report.md").read_text(encoding="utf-8")
            for heading in ("## 1. Tóm tắt", "## 3. So sánh backbone", "## 7. Kết luận", "## 8. Hạn chế", "## 9. Phụ lục"):
                self.assertIn(heading, report)
            self.assertIsNone(re.search(r"nan", report, re.IGNORECASE))  # no unformatted NaN


if __name__ == "__main__":
    unittest.main()
